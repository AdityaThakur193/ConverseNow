import os
import markdown
from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_inline_elements(paragraph, element):
    for child in element.children:
        if isinstance(child, NavigableString):
            run = paragraph.add_run(str(child))
        elif child.name == 'strong' or child.name == 'b':
            run = paragraph.add_run(child.text)
            run.bold = True
        elif child.name == 'em' or child.name == 'i':
            run = paragraph.add_run(child.text)
            run.italic = True
        elif child.name == 'code':
            run = paragraph.add_run(child.text)
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(199, 37, 78) # Ruby color for code
        elif child.name == 'a':
            run = paragraph.add_run(child.text)
            run.underline = True
            run.font.color.rgb = RGBColor(0, 0, 238)
        else:
            # Fallback for nested children
            add_inline_elements(paragraph, child)

def md_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    # Convert Markdown to HTML
    html = markdown.markdown(md_text, extensions=['extra', 'codehilite'])
    
    # Initialize DOCX
    doc = Document()
    
    # Adjust margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    soup = BeautifulSoup(html, 'html.parser')
    
    for element in soup.children:
        if isinstance(element, NavigableString):
            continue
            
        if element.name == 'p':
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            add_inline_elements(p, element)
            
        elif element.name and element.name.startswith('h') and len(element.name) == 2 and element.name[1].isdigit():
            level = int(element.name[1])
            # Add an extra empty paragraph for h1/h2 spacing
            if level <= 2:
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            p = doc.add_heading('', level=level)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            add_inline_elements(p, element)
            
        elif element.name == 'ul':
            for li in element.find_all('li', recursive=False):
                p = doc.add_paragraph(style='List Bullet')
                p.paragraph_format.space_after = Pt(4)
                add_inline_elements(p, li)
                
        elif element.name == 'ol':
            for li in element.find_all('li', recursive=False):
                p = doc.add_paragraph(style='List Number')
                p.paragraph_format.space_after = Pt(4)
                add_inline_elements(p, li)
                
        elif element.name == 'pre':
            # Code block / Mermaid diagram block
            code = element.find('code')
            text = code.text if code else element.text
            
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            
            run = p.add_run(text.strip())
            run.font.name = 'Consolas'
            run.font.size = Pt(9.0)
            run.font.color.rgb = RGBColor(51, 51, 51)
            
        elif element.name == 'hr':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run("❖   ❖   ❖").bold = True
            
    doc.save(docx_path)
    print(f"Successfully converted {md_path} to {docx_path}")

if __name__ == "__main__":
    md_to_docx("FINAL_PROJECT_REPORT.md", "FINAL_PROJECT_REPORT.docx")
