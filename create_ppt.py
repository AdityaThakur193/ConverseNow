import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def apply_text_styling(paragraph, font_name='Arial', font_size=16, bold=False, italic=False, color_rgb=(51,51,51)):
    paragraph.font.name = font_name
    paragraph.font.size = Pt(font_size)
    paragraph.font.bold = bold
    paragraph.font.italic = italic
    paragraph.font.color.rgb = RGBColor(*color_rgb)

def add_slide_with_bullets(prs, title_text, bullet_points):
    # Layout 1 is "Title and Content"
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    
    # Set Slide Title
    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.name = 'Arial'
    title.text_frame.paragraphs[0].font.size = Pt(36)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 51, 102) # Dark Blue
    
    # Set Content Placeholder
    content_placeholder = slide.placeholders[1]
    text_frame = content_placeholder.text_frame
    text_frame.clear() # Clear default
    
    for i, pt in enumerate(bullet_points):
        # Determine level and text
        if pt.startswith("  - "): # Sub-bullet
            level = 1
            text = pt[4:]
            bold = False
            size = 15
        elif pt.startswith("- "): # Main bullet
            level = 0
            text = pt[2:]
            bold = True
            size = 18
        else:
            level = 0
            text = pt
            bold = False
            size = 16
            
        p = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        p.text = text
        p.level = level
        p.space_after = Pt(8)
        apply_text_styling(p, font_name='Arial', font_size=size, bold=bold, color_rgb=(60,60,60))
        
    return slide

def main():
    prs = Presentation()
    
    # Set Slide Dimensions to 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # --- SLIDE 1: Title Slide ---
    slide_layout = prs.slide_layouts[0] # Title slide layout
    slide = prs.slides.add_slide(slide_layout)
    
    title = slide.shapes.title
    title.text = "ConverseNow"
    title.text_frame.paragraphs[0].font.size = Pt(54)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 51, 102)
    
    subtitle = slide.placeholders[1]
    subtitle.text = "Real-Time Assistive Translation and Indian Sign Language (ISL) Interpretation\n\nFaculty Capstone Presentation"
    subtitle.text_frame.paragraphs[0].font.size = Pt(22)
    subtitle.text_frame.paragraphs[0].font.color.rgb = RGBColor(102, 102, 102)
    
    # --- SLIDE 2: Introduction & Motivation ---
    add_slide_with_bullets(
        prs,
        "Introduction & Motivation",
        [
            "- The Communication Gap",
            "  - Signing individuals cannot easily communicate with the vocal majority in public kiosks, banks, and hospitals.",
            "- Multilingual Context in India",
            "  - Most sign translation platforms focus exclusively on English. Assistive systems in India must bridge regional scripts (Hindi, Telugu, Tamil, etc.).",
            "- Problem Statement",
            "  - Traditional translation pipelines require waiting for full audio clips to complete, causing unnatural conversation gaps.",
            "- Solution",
            "  - ConverseNow provides a stateless, real-time, low-latency WebSocket speech-to-sign pipeline driving a 3D humanoid avatar."
        ]
    )
    
    # --- SLIDE 3: Software Requirements Specification (SRS) ---
    add_slide_with_bullets(
        prs,
        "System & Functional Requirements (SRS)",
        [
            "- Input Requirements (FR-1 & FR-2)",
            "  - Capture microphone input at 16kHz Mono PCM16.",
            "  - Run real-time Speech Activity Detection (SAD) to segment utterances on 800ms of silence.",
            "- Translation Heuristics (FR-3 & FR-4)",
            "  - Perform STT in English and 10 regional Indian scripts via local Whisper model.",
            "  - Translate regional output to target subtitle text and base English translation.",
            "- Mapping & Rendering (FR-5 & FR-6)",
            "  - Map English base text into sign language gloss sequences.",
            "  - Stream glosses to Unity to trigger cross-fade humanoid clip animations in real time."
        ]
    )
    
    # --- SLIDE 4: System Architecture Overview ---
    add_slide_with_bullets(
        prs,
        "System Architecture & Data Flows",
        [
            "- Client-Server Architecture",
            "  - Unity Client (Frontend) <--> FastAPI WebSocket (Backend).",
            "- Stream Ingestion Router",
            "  - WS `/ws/audio` accepts raw 200ms audio bytes, runs local RMS energy checks, and segments speech on silence.",
            "- Concurrency & Event Loop Optimization",
            "  - Wrapped heavy tasks (Whisper, Translation, Gloss Mapping) in thread pools (`asyncio.to_thread`) to prevent blocking the async loop.",
            "- Payload Protocol",
            "  - Client sends target configuration message. Server responds with unified JSON payload containing transcriptions, target subtitles, and animation gloss lists."
        ]
    )
    
    # --- SLIDE 5: Speech-to-Text & Validation (STT) ---
    add_slide_with_bullets(
        prs,
        "Speech Activity Detection & Whisper STT",
        [
            "- RMS-Based Speech Segmentation",
            "  - Energy-based segmentation checks for voice threshold (cutoff: 400). Finalizes after 800ms of silence.",
            "- Prompt Steering for Regional Dialects",
            "  - Injects native initial script prompts into Whisper's decoder (e.g., Devanagari text for Hindi) to steer predicting tokens into clean regional orthography.",
            "- Unicode Script Validators",
            "  - Inspects code points of outputs against block ranges (e.g., Devanagari, Telugu, Tamil). Rejects ASCII/garbage output under regional translation modes."
        ]
    )
    
    # --- SLIDE 6: Translation & Cascade Gloss Engine ---
    add_slide_with_bullets(
        prs,
        "Translation & Cascade Gloss Translation",
        [
            "- External translation Integration",
            "  - Integrates SarvamAI API translation client (`sarvam-translate:v1`). Gracefully falls back to source text on connection or authorization errors.",
            "- Cascade Gloss Mapping Steps",
            "  - 1. Omit auxiliary stop-words ('is', 'to', 'the') carrying no semantic weight.",
            "  - 2. Map direct dictionary terms.",
            "  - 3. Substring match compound words (restricted to keys >= 3 characters to avoid short letter collisions).",
            "  - 4. Similarity match using SequenceMatcher ratio (strict cutoff: 0.75).",
            "  - 5. Fingerspelling fallback: spell out unrecognized words letter-by-letter using A-Z animations."
        ]
    )
    
    # --- SLIDE 7: Unity Client Playback & UI ---
    add_slide_with_bullets(
        prs,
        "Unity Frontend & Animation Playback",
        [
            "- Audio Ingestion",
            "  - `AudioStreamer.cs` captures mic inputs, converts floats to 16-bit PCM bytes, and sends them via WebSocket client.",
            "- Playback Manager (`ISLPlaybackManager.cs`)",
            "  - Builds length cache of loaded humanoid animations. Triggers animator cross-fades and handles timing transitions.",
            "- UI Subtitles (`SubtitleManager.cs`)",
            "  - Automatically links TMPro text components in the scene and displays translations.",
            "- Fingerspelling Sync",
            "  - Displays static 'Fingerspelled-<WORD>' labels to prevent rapid letter flicker on UI while the avatar signs spelling letters."
        ]
    )
    
    # --- SLIDE 8: Verification & Testing Report ---
    add_slide_with_bullets(
        prs,
        "Testing, Validation & Diagnostics",
        [
            "- Unit Testing Suite",
            "  - `test_gloss_generation.py` runs assertions on cleaning, tokenization, synonym mapping, and metrics. (100% Passed).",
            "  - `test_whisper_indic.py` validates initial prompt parameters, temperature models, and script ranges. (100% Passed).",
            "- Diagnostic Diagnostics",
            "  - Terminal outputs live volume levels (`Volume: {rms}`) to verify hardware streams during deployment.",
            "- Mock Tester Utility",
            "  - Context-menu inspector utility (`MockBackendTester.cs`) allows playing mock sequences offline without running python backend."
        ]
    )
    
    # --- SLIDE 9: Team Contributions & Roles ---
    add_slide_with_bullets(
        prs,
        "Team Contributions & Project Roles",
        [
            "- Aditya: Backend websocket server, SAD, Whisper STT prompts, Sarvam API, Gloss translation algorithms, async thread optimization, and tests.",
            "- Moulika: Unity project configuration, TMPro setup, UI layout, audio streaming connection client (`AudioStreamer.cs`), start/stop controllers, settings sliders.",
            "- Jaahnavi Neelam: ISL literature survey, humanoid rig animation integration, playback synchronization, bone rotation fixes, and MR implementation roadmap.",
            "- Sophie: MediaPipe coordinate translation pipelines, video-to-avatar mapping, Blender frame-by-frame joint edits, and pose tracking validations.",
            "- Shanmukhee: ISL animation dataset building, .anim clip extraction, avatar rig custom selections, validation tests, and AI animation tools benchmarking."
        ]
    )
    
    # Save Presentation
    prs.save("FINAL_PRESENTATION.pptx")
    print("Successfully created FINAL_PRESENTATION.pptx")

if __name__ == "__main__":
    main()
